from docx import Document

def parse_docx(file_path):
    doc = Document(file_path)
    text = ""
    metadata = {
        "source_filename": file_path.name,
        "paragraphs": len(doc.paragraphs)
    }
    for para in doc.paragraphs:
        if para.text.strip():
            text += para.text.strip() + "\n"
    return text.strip(), metadata
